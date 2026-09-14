"""Build a researcher's dataset: many variables, many days, one file.

## What this is for

A researcher does not want a raster. They want "sea-surface temperature and wind
speed off Kerala, from 1 June to today, as a spreadsheet" — several variables,
a date range, one file they can open.

Everything needed for that already existed in pieces: the ERDDAP grid fetcher
takes a ``when``, the catalogue knows which datasets serve which variables, and
the intent parser turns prose into a box and a range. This assembles them.

## Two decisions that shape the output

**A time series at one point, not a grid per day.** A 250x400 grid over 90 days
is nine million rows and no spreadsheet opens it. What a researcher almost always
wants from a range is *how a place changed*, so the builder samples a small set
of points — the box's centre by default, or a coarse lattice if asked — and walks
time. A grid for one day is what ``/research/export`` already does, and the two
are kept separate rather than merged into one endpoint with a mode flag.

**Every request is bounded before it runs.** Days times points times variables
is the cost, and it is computed and checked up front. A researcher who asks for
something too large gets a refusal that names the number and suggests the
reduction, rather than a request that dies at a proxy timeout with nothing to
show for it.

## The file carries its own provenance

The Excel workbook gets a second sheet — ``Provenance`` — with the dataset ids,
the endpoints, the licences, the caveats and a citation line. A spreadsheet that
leaves the system with no record of where its numbers came from is exactly how a
figure ends up in a paper with the wrong attribution.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from typing import TYPE_CHECKING, Any

from orca.provenance import utcnow
from orca.research.catalogue import BY_ID, CATALOGUE, CATALOGUE_VERSION, Dataset

if TYPE_CHECKING:
    from orca.research.variables import BuildVariable

log = logging.getLogger(__name__)

#: Hard ceiling on (days x points x variables). Roughly what fits comfortably in
#: a spreadsheet and in an upstream's patience.
MAX_CELLS = 20_000

#: Days a single build may span. Two years of daily data at a handful of points
#: is a legitimate research request; a decade is a bulk download and belongs at
#: the provider.
MAX_DAYS = 800

#: Points sampled across the box when the caller asks for more than the centre.
#: A 3x3 lattice shows a spatial gradient without turning a time series into a
#: raster.
LATTICE = 3


@dataclass(slots=True)
class BuildRequest:
    variables: list[str]
    west: float
    south: float
    east: float
    north: float
    start: datetime
    end: datetime
    #: 1 = the box centre only. Otherwise a LATTICE x LATTICE grid.
    points: int = 1
    step_days: int = 1


@dataclass(slots=True)
class BuildResult:
    rows: list[dict[str, Any]]
    variables: list[str]
    datasets: list[Dataset]
    requested: BuildRequest
    #: Variables asked for that no catalogue dataset can serve.
    unavailable: list[str] = field(default_factory=list)
    #: (day, variable) pairs the upstream had no value for.
    gaps: int = 0
    #: Missing count per variable. A single total says "25 values are missing";
    #: this says WHICH column is empty, which is the difference between a file a
    #: researcher can debug and one they quietly mistrust.
    gaps_by_variable: dict[str, int] = field(default_factory=dict)
    days_attempted: int = 0

    def empty_columns(self) -> dict[str, str]:
        """Variables with NO value at all, and the likely reason.

        A column of blanks is the most misleading thing a dataset file can
        contain: it reads as "measured and absent" rather than "never available
        at this cadence". The commonest cause here is asking a monthly composite
        for daily values, so that is named explicitly.
        """
        total = len(self.rows)
        if not total:
            return {}
        from orca.research.variables import BY_NAME as VARIABLE_REGISTRY

        explained: dict[str, str] = {}
        for variable in self.variables:
            if self.gaps_by_variable.get(variable, 0) < total:
                continue

            # A range variable has no catalogue Dataset behind it, so the
            # dataset-shaped reasoning below would report "unknown cadence" for
            # something whose provider and caveat we know exactly. Answer from
            # the registry instead.
            entry = VARIABLE_REGISTRY.get(variable)
            if entry is not None and entry.is_range_source:
                if (
                    entry.available_from
                    and self.requested.end.strftime("%Y-%m-%d") < entry.available_from
                ):
                    explained[variable] = (
                        f"{entry.provider} has no data before {entry.available_from}, and your "
                        "range ends before that."
                    )
                else:
                    explained[variable] = (
                        f"{entry.provider} returned no value for any requested day. The endpoint "
                        "answered, so this is coverage rather than a failure — the model may not "
                        f"resolve this point. {entry.caveat}"
                    )
                continue

            dataset = next(
                (d for d in self.datasets if any(v.name == variable for v in d.variables)), None
            )
            cadence = (dataset.cadence if dataset else "unknown").lower()
            if "month" in cadence:
                explained[variable] = (
                    f"{dataset.title if dataset else variable} is a MONTHLY composite, so it has "
                    "no value on an individual day. Request a monthly step, or use a daily "
                    "product for this variable."
                )
            else:
                # An archive whose coverage ended is the commonest cause, and it
                # is knowable rather than guessable: the ERDDAP registry records
                # the last timestamp each dataset actually contains.
                ends = _coverage_end(dataset)
                if ends is not None and self.requested.start.date() > ends.date():
                    explained[variable] = (
                        f"{dataset.title if dataset else variable} is an ARCHIVE that ends "
                        f"{ends:%Y-%m-%d}. Your range starts after that, so there is nothing to "
                        "return. Ask for a period before that date, or use a live product."
                    )
                else:
                    explained[variable] = (
                        f"No value was returned for any requested day ({cadence} cadence). The "
                        "upstream may not cover this box or this period."
                    )
        return explained

    def sources(self) -> list[dict[str, Any]]:
        """One provenance record per source that actually contributed a column.

        Assembled from the variables DELIVERED rather than from the datasets the
        builder happened to consult. Before this, a workbook whose wave, wind and
        rainfall columns all came from Open-Meteo cited only the ERDDAP dataset
        behind its SST column — so eleven of twelve columns travelled with the
        wrong attribution, which is precisely the failure the Provenance sheet
        exists to prevent.
        """
        from orca.research.variables import BY_NAME as VARIABLE_REGISTRY

        records: list[dict[str, Any]] = []
        seen: set[str] = set()

        for dataset in self.datasets:
            served = [
                name
                for name in self.variables
                if any(v.name == name for v in dataset.variables)
                and not (name in VARIABLE_REGISTRY and VARIABLE_REGISTRY[name].is_range_source)
            ]
            if not served:
                continue
            records.append(
                {
                    "id": dataset.id,
                    "title": dataset.title,
                    "provider": dataset.provider,
                    "licence": dataset.licence,
                    "endpoint": dataset.endpoint,
                    "caveats": dataset.caveats,
                    "columns": served,
                }
            )
            seen.add(dataset.endpoint)

        # Group the range variables by endpoint: eight marine columns are one
        # source, not eight, and listing them separately would bury the caveats
        # that differ.
        grouped: dict[str, list[Any]] = {}
        for name in self.variables:
            entry = VARIABLE_REGISTRY.get(name)
            if entry is not None and entry.is_range_source:
                grouped.setdefault(entry.endpoint, []).append(entry)

        for endpoint, entries in grouped.items():
            if endpoint in seen:
                continue
            first = entries[0]
            records.append(
                {
                    "id": first.kind,
                    "title": f"{first.provider} daily series",
                    "provider": first.provider,
                    "licence": first.licence,
                    "endpoint": endpoint,
                    # Each variable's own caveat, because they genuinely differ:
                    # wind is in km/h not knots, precipitation is a sum where the
                    # rest are maxima, and direction follows the meteorological
                    # convention. One blanket sentence would lose all three.
                    "caveats": " | ".join(f"{e.name}: {e.caveat}" for e in entries),
                    "columns": [e.name for e in entries],
                }
            )

        return records

    def summary(self) -> dict[str, Any]:
        return {
            "rows": len(self.rows),
            "variables": self.variables,
            "unavailable_variables": self.unavailable,
            "days_attempted": self.days_attempted,
            "missing_values": self.gaps,
            "missing_by_variable": self.gaps_by_variable,
            "empty_columns": self.empty_columns(),
            "datasets": self.sources(),
            "catalogue_version": CATALOGUE_VERSION,
            "generated_at": utcnow().isoformat(),
        }


def plan(request: BuildRequest) -> dict[str, Any]:
    """Cost a request before running it, and say what it would cost.

    Returned rather than raised so the caller can show the number and offer the
    reduction. A refusal that names the figure is actionable; a timeout is not.
    """
    from orca.research.variables import BY_NAME as VARIABLE_REGISTRY

    days = max(1, (request.end - request.start).days // max(request.step_days, 1) + 1)
    points = 1 if request.points <= 1 else LATTICE * LATTICE
    cells = days * points * max(len(request.variables), 1)

    # Upstream requests are counted separately from cells, because the two no
    # longer track each other. A range variable costs ONE request per point
    # whatever the span; a griddap variable costs one per point per day. Costing
    # them the same would refuse a 90-day wave series that is nine calls.
    ranged = [
        name
        for name in request.variables
        if name in VARIABLE_REGISTRY and VARIABLE_REGISTRY[name].is_range_source
    ]
    per_day = [name for name in request.variables if name not in ranged]
    range_kinds = {VARIABLE_REGISTRY[name].kind for name in ranged}
    requests = points * (len(range_kinds) + days * len(per_day))

    return {
        "days": days,
        "points": points,
        "variables": len(request.variables),
        "cells": cells,
        "upstream_requests": requests,
        "range_variables": sorted(ranged),
        "per_day_variables": sorted(per_day),
        "within_limits": days <= MAX_DAYS and cells <= MAX_CELLS,
        "max_cells": MAX_CELLS,
        "max_days": MAX_DAYS,
        "cost_note": (
            f"{len(ranged)} variable(s) answer the whole span in one request each, and "
            f"{len(per_day)} need one request per day. {requests} upstream request(s) in total."
        ),
    }


def _sources_for(variables: list[str]) -> tuple[dict[str, Dataset], list[str]]:
    """Pick one fetchable dataset per variable, preferring the finest resolution.

    Deterministic, like the rest of the matching: no model chooses a source for a
    researcher's file.

    "Fetchable" is stricter than "servable". The builder walks a date range
    through ERDDAP's griddap, so a dataset is only a candidate if it has an
    ERDDAP key here. Ranking on resolution alone picked Open-Meteo for wind —
    finest at 0.1 deg, and a forecast API with no griddap endpoint — so every
    wind cell came back empty while the file looked like it had simply found no
    data.
    """
    from orca.research.variables import BY_NAME as VARIABLE_REGISTRY

    chosen: dict[str, Dataset] = {}
    missing: list[str] = []
    for variable in variables:
        # A range variable has exactly one source, and it is not griddap. Without
        # this, `wind_speed` resolved to the ASCAT catalogue entry as well as to
        # ERA5 — and ASCAT is an archive that stopped in May 2023, so the answer
        # this function gave was one that could never return a value. `build()`
        # already routed it correctly; the trap was for whoever read this next.
        entry = VARIABLE_REGISTRY.get(variable)
        if entry is not None and entry.is_range_source:
            continue
        candidates = [
            d
            for d in CATALOGUE
            if d.servable
            and _erddap_key(d.id) is not None
            and any(v.name == variable for v in d.variables)
        ]
        if not candidates:
            missing.append(variable)
            continue
        chosen[variable] = min(candidates, key=lambda d: d.resolution_deg)
    return chosen, missing


def _coverage_end(dataset: Dataset | None) -> datetime | None:
    """The last timestamp a dataset actually contains, from the ERDDAP registry.

    Measured by the registry rather than assumed from the catalogue prose, so an
    archive that quietly stopped updating cannot masquerade as a live feed.
    """
    if dataset is None:
        return None
    from orca.sources.erddap import DATASETS

    key = _erddap_key(dataset.id)
    entry = DATASETS.get(key) if key else None
    return getattr(entry, "coverage_end", None)


def _serves_daily(dataset: Dataset) -> bool:
    """Whether asking this dataset for one specific day can ever work.

    A monthly composite cannot. Checking up front matters because ORCA's HTTP
    client retries four times on failure — entirely correctly, for a transient
    error — so twenty daily requests against a monthly product became eighty
    calls to somebody else's server for an answer that was never going to exist.
    """
    return "month" not in dataset.cadence.lower()


def _points(request: BuildRequest) -> list[tuple[float, float, str]]:
    if request.points <= 1:
        lat = (request.south + request.north) / 2
        lon = (request.west + request.east) / 2
        return [(lat, lon, "centre")]
    out: list[tuple[float, float, str]] = []
    for row in range(LATTICE):
        for col in range(LATTICE):
            lat = request.south + (request.north - request.south) * (row + 0.5) / LATTICE
            lon = request.west + (request.east - request.west) * (col + 0.5) / LATTICE
            out.append((lat, lon, f"r{row}c{col}"))
    return out


#: Concurrent upstream fetches. Bounded rather than unlimited: ERDDAP servers are
#: run by research institutions on modest hardware, and a builder that opens
#: sixty simultaneous connections to one of them is the reason free endpoints get
#: rate limits. Eight is fast enough to make a 30-day build interactive and
#: polite enough not to be noticed.
CONCURRENCY = 8

#: Open-Meteo answers a whole date span in one request. Ninety days of wave
#: height is ONE call, not ninety — which is why the range sources are fetched
#: per (point, endpoint) and the ERDDAP ones per (point, day).
_OPEN_METEO_TIMEOUT_S = 45.0


def _range_endpoint(kind: str) -> str:
    from orca.research.variables import _ERA5_URL, _MARINE_URL

    return _MARINE_URL if kind == "openmeteo_marine" else _ERA5_URL


async def _fetch_range(
    *,
    kind: str,
    variables: list[BuildVariable],
    lat: float,
    lon: float,
    start: datetime,
    end: datetime,
) -> dict[str, dict[str, float | None]]:
    """One request covering the whole span, keyed {variable: {date: value}}.

    Returns an empty mapping on any failure rather than raising. A researcher
    asking for eight variables should not lose the other seven because the
    marine endpoint had a bad minute — and `gaps_by_variable` will show exactly
    which column came back empty.
    """
    from orca.sources.base import get_client

    params = {
        "latitude": lat,
        "longitude": lon,
        "daily": ",".join(v.api_name for v in variables),
        "start_date": start.strftime("%Y-%m-%d"),
        "end_date": end.strftime("%Y-%m-%d"),
        "timezone": "UTC",
    }
    try:
        client = await get_client()
        response = await client.get(
            _range_endpoint(kind), params=params, timeout=_OPEN_METEO_TIMEOUT_S
        )
        response.raise_for_status()
        daily = response.json().get("daily") or {}
    except Exception as exc:  # noqa: BLE001
        log.warning("builder: %s range fetch failed at %.2f,%.2f: %s", kind, lat, lon, exc)
        return {}

    dates = daily.get("time") or []
    out: dict[str, dict[str, float | None]] = {}
    for variable in variables:
        series = daily.get(variable.api_name) or []
        out[variable.name] = {
            day: (round(float(value), 4) if isinstance(value, (int, float)) else None)
            for day, value in zip(dates, series)
        }
    return out


async def build(request: BuildRequest) -> BuildResult:
    """Walk the date range and assemble a tidy table.

    One row per (day, point), one column per variable — "tidy" in the sense a
    statistician means, because that is the shape every analysis package expects
    and reshaping a wide table is a chore nobody thanks you for.

    Fetches run CONCURRENTLY under a semaphore. Sequentially, a 20-day build of
    three variables is sixty round trips at a second or more each, which is a
    minute of a researcher staring at a spinner and long enough for a proxy to
    give up on the request entirely.
    """
    import asyncio

    from orca.research.variables import BY_NAME as VARIABLE_REGISTRY
    from orca.sources import erddap

    # Two dispatch paths, because the sources answer different questions.
    # A range source returns a whole span in one response; ERDDAP's griddap
    # answers for one timestamp. Treating them identically is what made the
    # first builder slow AND narrow — it looped days and issued a request per
    # day per variable regardless, so a range source was hit ninety times for
    # data it would have returned once, and anything without a griddap endpoint
    # was simply unreachable.
    requested = list(dict.fromkeys(request.variables))
    range_vars = [
        VARIABLE_REGISTRY[name]
        for name in requested
        if name in VARIABLE_REGISTRY and VARIABLE_REGISTRY[name].is_range_source
    ]
    erddap_names = [
        name
        for name in requested
        if name not in VARIABLE_REGISTRY or not VARIABLE_REGISTRY[name].is_range_source
    ]

    sources, missing = _sources_for(erddap_names)
    served_erddap = [v for v in erddap_names if v in sources]
    # A variable is only "unavailable" if NEITHER path can serve it.
    missing = [name for name in missing if name not in VARIABLE_REGISTRY]
    served = [
        name
        for name in requested
        if name in served_erddap or any(v.name == name for v in range_vars)
    ]
    points = _points(request)

    days: list[datetime] = []
    current = request.start
    while current <= request.end:
        days.append(current)
        current += timedelta(days=max(request.step_days, 1))

    limit = asyncio.Semaphore(CONCURRENCY)

    async def fetch(day: datetime, lat: float, lon: float, variable: str) -> float | None:
        dataset = sources[variable]
        key = _erddap_key(dataset.id)
        if key is None or not _serves_daily(dataset):
            # Not attempted: a monthly composite has no value for one day, and
            # the retry policy would turn each impossible request into four.
            return None
        async with limit:
            try:
                found = await erddap.erddap.point(key, lat, lon, when=day, variables=[variable])
            except Exception as exc:  # noqa: BLE001 — one bad day is not fatal
                log.debug("builder: %s %s failed (%s)", variable, day.date(), exc)
                return None
        evidence = found.get(variable)
        if evidence is not None and isinstance(evidence.value, (int, float)):
            return round(float(evidence.value), 4)
        return None

    jobs = [
        (day, lat, lon, label, variable)
        for day in days
        for lat, lon, label in points
        for variable in served_erddap
    ]
    values = await asyncio.gather(
        *(fetch(day, lat, lon, variable) for day, lat, lon, _, variable in jobs)
    )

    # The range sources: one request per (point, endpoint), not per day. Eight
    # marine variables over ninety days at nine points is 9 calls, against the
    # 6,480 the per-day path would have issued.
    by_kind: dict[str, list[Any]] = {}
    for variable in range_vars:
        by_kind.setdefault(variable.kind, []).append(variable)

    async def range_job(kind: str, group: list[Any], lat: float, lon: float, label: str):
        async with limit:
            series = await _fetch_range(
                kind=kind,
                variables=group,
                lat=lat,
                lon=lon,
                start=request.start,
                end=request.end,
            )
        return label, series

    range_results = await asyncio.gather(
        *(
            range_job(kind, group, lat, lon, label)
            for kind, group in by_kind.items()
            for lat, lon, label in points
        )
    )
    #: {point label: {variable: {date: value}}}
    ranged: dict[str, dict[str, dict[str, float | None]]] = {}
    for label, series in range_results:
        ranged.setdefault(label, {}).update(series)

    table: dict[tuple[str, str], dict[str, Any]] = {}
    per_variable: dict[str, int] = {}
    gaps = 0
    for (day, lat, lon, label, variable), value in zip(jobs, values, strict=True):
        key = (day.strftime("%Y-%m-%d"), label)
        row = table.setdefault(
            key,
            {
                "date": key[0],
                "latitude": round(lat, 4),
                "longitude": round(lon, 4),
                "point": label,
            },
        )
        row[variable] = value
        if value is None:
            gaps += 1
            per_variable[variable] = per_variable.get(variable, 0) + 1

    # Merge the range series in. Every (day, point) cell is written even when the
    # provider had no value, so a gap is recorded as a gap rather than as a
    # missing key that would silently widen into a ragged table.
    if range_vars:
        for day in days:
            date_key = day.strftime("%Y-%m-%d")
            for lat, lon, label in points:
                row = table.setdefault(
                    (date_key, label),
                    {
                        "date": date_key,
                        "latitude": round(lat, 4),
                        "longitude": round(lon, 4),
                        "point": label,
                    },
                )
                for variable in range_vars:
                    value = ranged.get(label, {}).get(variable.name, {}).get(date_key)
                    row[variable.name] = value
                    if value is None:
                        gaps += 1
                        per_variable[variable.name] = per_variable.get(variable.name, 0) + 1

    rows = [table[k] for k in sorted(table)]
    return BuildResult(
        rows=rows,
        variables=served,
        datasets=list({sources[v].id: sources[v] for v in served_erddap}.values()),
        requested=request,
        unavailable=missing,
        gaps=gaps,
        gaps_by_variable=per_variable,
        days_attempted=len(days),
    )


def _erddap_key(dataset_id: str) -> str | None:
    return {
        "mur_sst": "mur_sst",
        "esacci_chl_monthly": "esacci_chl_monthly",
        "incois_tmi_sst": "incois_tmi_sst",
        "incois_oceansat2": "incois_oceansat2",
        # The registry key, not the catalogue id. These diverged and the lookup
        # returned None silently, so every wind cell came back empty and the
        # "archive has ended" explanation never fired either — it depends on the
        # same lookup.
        "ascat_winds": "incois_ascat_wind",
    }.get(dataset_id)


# ----------------------------------------------------------------- rendering


def to_csv(result: BuildResult) -> str:
    """Tidy CSV with a commented provenance preamble."""
    import csv

    buffer = io.StringIO()
    meta = result.summary()
    buffer.write(f"# ORCA dataset build · {meta['generated_at']}\n")
    buffer.write(f"# catalogue: {meta['catalogue_version']}\n")
    for dataset in meta["datasets"]:
        buffer.write(
            f"# source: {dataset['title']} ({dataset['provider']}) — {dataset['licence']}\n"
        )
        buffer.write(f"#   endpoint: {dataset['endpoint']}\n")
        buffer.write(f"#   caveat: {dataset['caveats']}\n")
    if result.unavailable:
        buffer.write(
            f"# NOT SERVED (no catalogue dataset provides these): {', '.join(result.unavailable)}\n"
        )
    buffer.write(f"# missing values: {result.gaps} — blank cells are upstream gaps, not zeroes\n")

    columns = ["date", "latitude", "longitude", "point", *result.variables]
    writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(result.rows)
    return buffer.getvalue()


def to_xlsx(result: BuildResult) -> bytes:
    """An .xlsx with the data on one sheet and its provenance on another.

    The second sheet is the point. A spreadsheet that leaves the system with no
    record of where its numbers came from is how a figure ends up in a paper with
    the wrong attribution — and a commented CSV header does not survive the trip
    through Excel, which is exactly where this file is going.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    meta = result.summary()
    book = Workbook()

    sheet = book.active
    sheet.title = "Data"
    columns = ["date", "latitude", "longitude", "point", *result.variables]
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in result.rows:
        sheet.append([row.get(name) for name in columns])
    for index, name in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(11, len(name) + 3)

    notes = book.create_sheet("Provenance")
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 110
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    def line(label: str, value: str) -> None:
        notes.append([label, value])
        notes.cell(row=notes.max_row, column=1).font = bold
        notes.cell(row=notes.max_row, column=2).alignment = wrap

    line("Generated by", "ORCA — Marine EcOsystem Reasoning with Collaborative Agents")
    line("Generated at", meta["generated_at"])
    line("Catalogue version", meta["catalogue_version"])
    line(
        "Area",
        f"{result.requested.west}, {result.requested.south} to "
        f"{result.requested.east}, {result.requested.north}",
    )
    line("Period", f"{result.requested.start:%Y-%m-%d} to {result.requested.end:%Y-%m-%d}")
    line("Rows", str(len(result.rows)))
    line("Missing values", f"{result.gaps} — blank cells are upstream gaps, not zeroes")
    for variable, why in result.empty_columns().items():
        line(f"EMPTY: {variable}", why)
    if result.unavailable:
        line(
            "NOT served",
            ", ".join(result.unavailable) + " — no dataset in ORCA's catalogue provides these",
        )
    notes.append([])

    for dataset in meta["datasets"]:
        line("Source", dataset["title"])
        # Which columns came from this source, named. A researcher merging two
        # ORCA files needs to know that `wind_speed` is ERA5 reanalysis and
        # `sst` is a satellite analysis; a list of sources without that mapping
        # leaves them to guess.
        line("  columns", ", ".join(dataset.get("columns") or []))
        line("  provider", dataset["provider"])
        line("  licence", dataset["licence"])
        line("  endpoint", dataset["endpoint"])
        line("  caveat", dataset["caveats"])
        notes.append([])

    line(
        "How to cite",
        "; ".join(
            f"{d['provider']}. {d['title']}. Accessed via ORCA on "
            f"{utcnow():%Y-%m-%d}. {d['licence']}."
            for d in meta["datasets"]
        ),
    )

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def parse_day(value: str, *, fallback: datetime) -> datetime:
    """ISO date, or 'today'/'now'. Researchers write both."""
    text = (value or "").strip().lower()
    if text in ("", "today", "now", "current"):
        return fallback
    try:
        return datetime.strptime(text, "%Y-%m-%d").replace(hour=9, tzinfo=UTC)
    except ValueError as exc:
        raise ValueError(f"expected YYYY-MM-DD or 'today', got {value!r}") from exc


def known_variables() -> list[str]:
    """Variables the builder can ACTUALLY deliver, not merely ones we list.

    A variable only qualifies if some servable dataset both has an ERDDAP key
    here AND declares a mapping for it in the registry. `wind_speed` failed that
    second test: the ASCAT entry carries no variable mapping, so a request for it
    returned an empty column on every row. Advertising a variable and then
    handing back blanks is worse than not offering it — the researcher cannot
    tell "no data today" from "never possible".
    """
    from orca.research.variables import BY_NAME as VARIABLE_REGISTRY
    from orca.sources.erddap import DATASETS

    available: set[str] = set()
    for dataset in CATALOGUE:
        if not dataset.servable:
            continue
        key = _erddap_key(dataset.id)
        entry = DATASETS.get(key) if key else None
        if entry is None:
            continue
        mapped = set(entry.variables.values())
        for variable in dataset.variables:
            if variable.name in mapped:
                available.add(variable.name)

    # Plus everything in the verified registry. The rule there is the inverse of
    # the catalogue's: a variable appears only if it has been fetched
    # successfully at least once, so the list is short and true rather than long
    # and aspirational.
    available.update(VARIABLE_REGISTRY)
    return sorted(available)


__all__ = [
    "BY_ID",
    "BuildRequest",
    "BuildResult",
    "build",
    "known_variables",
    "parse_day",
    "plan",
    "to_csv",
    "to_xlsx",
]
