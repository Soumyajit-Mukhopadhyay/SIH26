"""The dataset builder: what it promises, and what it refuses to promise.

Most of these pin honesty properties rather than arithmetic. A dataset file is
consumed hours later by somebody who cannot see the system that produced it, so
the failure that matters is not a crash — it is a column of blanks that reads as
"measured and absent" when it actually means "never possible".
"""

from __future__ import annotations

from datetime import UTC, datetime

import pytest

from orca.research import builder


def spec(**overrides: object) -> builder.BuildRequest:
    base = {
        "variables": ["sst"],
        "west": 74.0,
        "south": 8.0,
        "east": 77.5,
        "north": 13.0,
        "start": datetime(2026, 8, 1, 9, tzinfo=UTC),
        "end": datetime(2026, 8, 10, 9, tzinfo=UTC),
    }
    base.update(overrides)
    return builder.BuildRequest(**base)  # type: ignore[arg-type]


class TestAdvertisedCapability:
    def test_only_variables_that_can_actually_be_delivered_are_offered(self) -> None:
        """Every advertised variable must be reachable by one of the two paths,
        and by exactly the path it claims.

        The rule used to be "some catalogue dataset has an ERDDAP key AND a
        registry variable mapping". That is now only one of the two ways a
        column gets filled: a range variable is served straight from
        `research.variables`, which earns its place by having been fetched
        successfully at least once rather than by appearing in a catalogue.

        What has NOT changed is the reason the test exists. `wind_speed` was
        advertised off a catalogue entry that carried no variable mapping and
        returned an empty column on every row, and a researcher cannot tell "no
        data today" from "never possible"."""
        from orca.research.variables import BY_NAME as REGISTRY
        from orca.sources.erddap import DATASETS

        for variable in builder.known_variables():
            entry = REGISTRY.get(variable)
            if entry is not None and entry.is_range_source:
                # Served by a range endpoint. It needs no catalogue dataset, but
                # it does need everything the Provenance sheet will quote.
                assert entry.endpoint.startswith("https://")
                assert entry.provider and entry.licence and entry.caveat
                continue

            sources, missing = builder._sources_for([variable])
            assert not missing, f"{variable} is advertised but has no source"
            dataset = sources[variable]
            key = builder._erddap_key(dataset.id)
            assert key in DATASETS, f"{variable} maps to a non-existent registry key {key!r}"
            assert variable in set(DATASETS[key].variables.values())

    def test_a_range_variable_is_never_also_routed_through_griddap(self) -> None:
        """The two paths must not both claim a variable.

        If they did, the per-day path would issue one request per day for data
        the range path already has — and the two answers would differ, because
        they are different products. `sea_surface_temperature` (Open-Meteo's
        model field) and `sst` (NASA MUR satellite analysis) are deliberately
        separate names for exactly this reason."""
        from orca.research.variables import BY_NAME as REGISTRY

        for name, entry in REGISTRY.items():
            if not entry.is_range_source:
                continue
            sources, _ = builder._sources_for([name])
            assert name not in sources, (
                f"{name} is served by both the range endpoint and griddap; "
                "one of the two must be renamed"
            )

    def test_every_catalogue_key_resolves_to_a_real_registry_entry(self) -> None:
        """The catalogue id and the registry key diverged once — `ascat_winds`
        against `incois_ascat_wind` — and the lookup returned None silently, so
        the data was empty AND the explanation for why never fired."""
        from orca.research.catalogue import CATALOGUE
        from orca.sources.erddap import DATASETS

        broken = [
            (d.id, builder._erddap_key(d.id))
            for d in CATALOGUE
            if builder._erddap_key(d.id) is not None and builder._erddap_key(d.id) not in DATASETS
        ]
        assert broken == [], f"catalogue ids mapping to unknown registry keys: {broken}"

    def test_a_source_is_only_chosen_if_the_builder_can_fetch_it(self) -> None:
        """Ranking on resolution alone picked Open-Meteo for wind — finest at
        0.1 deg, and a forecast API with no griddap endpoint."""
        sources, _ = builder._sources_for(["sst"])
        assert builder._erddap_key(sources["sst"].id) is not None


class TestCostControl:
    def test_a_request_is_costed_before_it_runs(self) -> None:
        plan = builder.plan(spec(variables=["sst", "chlorophyll"], points=9))
        assert plan["cells"] == plan["days"] * plan["points"] * plan["variables"]
        assert plan["points"] == builder.LATTICE * builder.LATTICE

    def test_an_oversized_request_is_refused_with_the_figure(self) -> None:
        """A refusal that names the number is actionable; a proxy timeout is not."""
        huge = spec(
            start=datetime(2020, 1, 1, 9, tzinfo=UTC), points=9, variables=["sst", "chlorophyll"]
        )
        plan = builder.plan(huge)
        assert plan["within_limits"] is False
        assert plan["cells"] > builder.MAX_CELLS

    def test_a_range_variable_costs_one_request_not_one_per_day(self) -> None:
        """The whole reason for the second path. Ninety days of wave height is
        ONE call; costing it as ninety would refuse a request that is trivially
        affordable."""
        long_span = spec(
            variables=["wave_height"],
            start=datetime(2026, 1, 1, 9, tzinfo=UTC),
            end=datetime(2026, 3, 31, 9, tzinfo=UTC),
        )
        costed = builder.plan(long_span)
        assert costed["days"] > 80
        assert costed["upstream_requests"] == 1
        assert costed["range_variables"] == ["wave_height"]

    def test_a_per_day_variable_still_costs_one_request_per_day(self) -> None:
        costed = builder.plan(spec(variables=["sst"]))
        assert costed["upstream_requests"] == costed["days"]
        assert costed["per_day_variables"] == ["sst"]

    def test_step_days_reduces_the_cost(self) -> None:
        dense = builder.plan(spec(start=datetime(2026, 1, 1, 9, tzinfo=UTC)))
        sparse = builder.plan(spec(start=datetime(2026, 1, 1, 9, tzinfo=UTC), step_days=7))
        assert sparse["cells"] < dense["cells"]


class TestImpossibleRequestsAreNotAttempted:
    def test_a_monthly_product_is_not_asked_for_a_single_day(self) -> None:
        """ORCA's HTTP client retries four times on failure — correctly, for a
        transient error. Twenty daily requests against a monthly composite became
        eighty calls to somebody else's server for an answer that could not exist.
        """
        from orca.research.catalogue import BY_ID

        assert builder._serves_daily(BY_ID["mur_sst"]) is True
        assert builder._serves_daily(BY_ID["esacci_chl_monthly"]) is False


class TestEmptyColumnsAreExplained:
    def _result(self, variable: str, dataset_id: str, rows: int = 3) -> builder.BuildResult:
        from orca.research.catalogue import BY_ID

        return builder.BuildResult(
            rows=[{"date": f"2026-08-0{i + 1}", variable: None} for i in range(rows)],
            variables=[variable],
            datasets=[BY_ID[dataset_id]],
            requested=spec(variables=[variable]),
            gaps=rows,
            gaps_by_variable={variable: rows},
            days_attempted=rows,
        )

    def test_a_monthly_composite_says_so(self) -> None:
        explained = self._result("chlorophyll", "esacci_chl_monthly").empty_columns()
        assert "MONTHLY" in explained["chlorophyll"]

    def test_an_expired_archive_names_its_end_date(self) -> None:
        """Uses `wind_direction`, not `wind_speed`.

        `wind_speed` used to come off this dead ASCAT archive and now comes from
        ERA5, so it no longer exercises this path at all — but the path still
        matters, because `wind_direction` is still only available from the
        archive and a request for it in 2026 returns nothing."""
        from orca.research.catalogue import BY_ID

        result = builder.BuildResult(
            rows=[{"date": "2026-08-01", "wind_direction": None}],
            variables=["wind_direction"],
            datasets=[BY_ID["ascat_winds"]],
            requested=spec(variables=["wind_direction"]),
            gaps=1,
            gaps_by_variable={"wind_direction": 1},
            days_attempted=1,
        )
        explained = result.empty_columns()
        assert "ARCHIVE" in explained["wind_direction"]
        assert "2023-05-21" in explained["wind_direction"]

    def test_a_range_variable_blames_its_provider_not_a_cadence(self) -> None:
        """A range variable has no catalogue Dataset behind it, so the
        dataset-shaped reasoning would have reported "unknown cadence" for
        something whose provider and caveat are known exactly."""
        result = builder.BuildResult(
            rows=[{"date": "2026-08-01", "wave_height": None}],
            variables=["wave_height"],
            datasets=[],
            requested=spec(variables=["wave_height"]),
            gaps=1,
            gaps_by_variable={"wave_height": 1},
            days_attempted=1,
        )
        explained = result.empty_columns()
        assert "unknown" not in explained["wave_height"]
        assert "Open-Meteo" in explained["wave_height"]

    def test_a_partially_full_column_is_not_flagged(self) -> None:
        """Only a column with NO value at all is explained. A few gaps are
        ordinary cloud cover and saying so on every file would be noise."""
        from orca.research.catalogue import BY_ID

        result = builder.BuildResult(
            rows=[{"sst": 29.1}, {"sst": None}, {"sst": 28.8}],
            variables=["sst"],
            datasets=[BY_ID["mur_sst"]],
            requested=spec(),
            gaps=1,
            gaps_by_variable={"sst": 1},
            days_attempted=3,
        )
        assert result.empty_columns() == {}


class TestRendering:
    def _result(self) -> builder.BuildResult:
        from orca.research.catalogue import BY_ID

        return builder.BuildResult(
            rows=[
                {
                    "date": "2026-08-01",
                    "latitude": 10.5,
                    "longitude": 75.7,
                    "point": "centre",
                    "sst": 29.1,
                },
                {
                    "date": "2026-08-02",
                    "latitude": 10.5,
                    "longitude": 75.7,
                    "point": "centre",
                    "sst": 29.3,
                },
            ],
            variables=["sst"],
            datasets=[BY_ID["mur_sst"]],
            requested=spec(),
            days_attempted=2,
        )

    def test_every_delivered_column_is_attributed(self) -> None:
        """A workbook whose wave, wind and rainfall columns all came from
        Open-Meteo once cited only the ERDDAP dataset behind its SST column, so
        eleven of twelve columns travelled with the wrong attribution — which is
        precisely what the Provenance sheet exists to prevent."""
        from orca.research.catalogue import BY_ID

        result = builder.BuildResult(
            rows=[{"date": "2026-08-01", "sst": 29.1, "wave_height": 1.2, "wind_speed": 14.0}],
            variables=["sst", "wave_height", "wind_speed"],
            datasets=[BY_ID["mur_sst"]],
            requested=spec(variables=["sst", "wave_height", "wind_speed"]),
            days_attempted=1,
        )
        attributed = {column for source in result.sources() for column in source["columns"]}
        assert attributed == {"sst", "wave_height", "wind_speed"}
        providers = {source["provider"] for source in result.sources()}
        assert len(providers) >= 2, "the Open-Meteo columns must not inherit NASA's attribution"

    def test_csv_carries_provenance_in_its_preamble(self) -> None:
        text = builder.to_csv(self._result())
        assert "# source:" in text
        assert "endpoint:" in text
        assert "caveat:" in text
        assert "date,latitude,longitude,point,sst" in text

    def test_xlsx_has_a_separate_provenance_sheet(self) -> None:
        """A commented CSV header does not survive a trip through Excel, which is
        exactly where this file is going."""
        import io as _io

        from openpyxl import load_workbook

        book = load_workbook(_io.BytesIO(builder.to_xlsx(self._result())))
        assert book.sheetnames == ["Data", "Provenance"]
        labels = [row[0] for row in book["Provenance"].iter_rows(values_only=True) if row[0]]
        assert "How to cite" in labels
        assert any("licence" in str(label) for label in labels)

    def test_the_data_sheet_is_tidy(self) -> None:
        """One row per (day, point), one column per variable — the shape every
        analysis package expects."""
        import io as _io

        from openpyxl import load_workbook

        sheet = load_workbook(_io.BytesIO(builder.to_xlsx(self._result())))["Data"]
        assert [c.value for c in sheet[1]] == ["date", "latitude", "longitude", "point", "sst"]
        assert sheet.max_row == 3


class TestDateParsing:
    def test_today_is_accepted_because_researchers_write_it(self) -> None:
        fallback = datetime(2026, 9, 13, 9, tzinfo=UTC)
        for text in ("today", "now", "current", ""):
            assert builder.parse_day(text, fallback=fallback) == fallback

    def test_an_iso_date_is_parsed(self) -> None:
        parsed = builder.parse_day("2025-03-04", fallback=datetime.now(UTC))
        assert (parsed.year, parsed.month, parsed.day) == (2025, 3, 4)

    def test_a_bad_date_is_rejected_with_the_value(self) -> None:
        with pytest.raises(ValueError, match="04-03-2025"):
            builder.parse_day("04-03-2025", fallback=datetime.now(UTC))
